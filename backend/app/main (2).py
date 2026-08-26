from datetime import datetime, timedelta, date, time
from pathlib import Path
import csv, io, os, secrets, hashlib
from typing import Optional
from fastapi import FastAPI, Depends, HTTPException, UploadFile, File, Form, Query
from fastapi.middleware.cors import CORSMiddleware
from fastapi.security import OAuth2PasswordBearer, OAuth2PasswordRequestForm
from fastapi.staticfiles import StaticFiles
from jose import jwt, JWTError
from passlib.context import CryptContext
from pydantic import BaseModel
from sqlalchemy import create_engine, String, Integer, Date, DateTime, Boolean, Float, ForeignKey, Text, UniqueConstraint, func, or_
from sqlalchemy.orm import DeclarativeBase, Mapped, mapped_column, sessionmaker, Session

DATABASE_URL=os.getenv('DATABASE_URL','sqlite:///./rizvi360.db'); SECRET=os.getenv('JWT_SECRET','dev-only-change-me'); ALGORITHM='HS256'
UPLOAD_DIR=Path(os.getenv('UPLOAD_DIR','./uploads')); UPLOAD_DIR.mkdir(parents=True,exist_ok=True)
engine=create_engine(DATABASE_URL,connect_args={'check_same_thread':False} if DATABASE_URL.startswith('sqlite') else {})
SessionLocal=sessionmaker(bind=engine,autocommit=False,autoflush=False)
pwd=CryptContext(schemes=['bcrypt'],deprecated='auto'); oauth2=OAuth2PasswordBearer(tokenUrl='/api/auth/login')
class Base(DeclarativeBase): pass
class Employee(Base):
    __tablename__='employees'; id:Mapped[int]=mapped_column(primary_key=True); employee_id:Mapped[str]=mapped_column(String(64),unique=True,index=True); card_no:Mapped[str]=mapped_column(String(64),unique=True,index=True); name:Mapped[str]=mapped_column(String(160)); designation:Mapped[Optional[str]]=mapped_column(String(160)); department:Mapped[Optional[str]]=mapped_column(String(160)); section:Mapped[Optional[str]]=mapped_column(String(160)); floor:Mapped[Optional[str]]=mapped_column(String(80)); line:Mapped[Optional[str]]=mapped_column(String(80)); join_date:Mapped[Optional[date]]=mapped_column(Date); mobile:Mapped[Optional[str]]=mapped_column(String(40)); email:Mapped[Optional[str]]=mapped_column(String(160)); grade:Mapped[Optional[str]]=mapped_column(String(40)); category:Mapped[Optional[str]]=mapped_column(String(50)); manager_id:Mapped[Optional[str]]=mapped_column(String(64)); status:Mapped[str]=mapped_column(String(30),default='ACTIVE',index=True); role:Mapped[str]=mapped_column(String(40),default='EMPLOYEE'); password_hash:Mapped[str]=mapped_column(String(255),default='')
class Punch(Base):
    __tablename__='punches'; id:Mapped[int]=mapped_column(primary_key=True); employee_id:Mapped[str]=mapped_column(String(64),index=True); punch_date:Mapped[date]=mapped_column(Date,index=True); in_time:Mapped[Optional[str]]=mapped_column(String(8)); out_time:Mapped[Optional[str]]=mapped_column(String(8)); break_minutes:Mapped[int]=mapped_column(Integer,default=60); source_file:Mapped[Optional[str]]=mapped_column(String(255)); __table_args__=(UniqueConstraint('employee_id','punch_date',name='uq_punch_day'),)
class Salary(Base):
    __tablename__='salaries'; id:Mapped[int]=mapped_column(primary_key=True); employee_id:Mapped[str]=mapped_column(String(64),index=True); salary_month:Mapped[str]=mapped_column(String(7),index=True); basic:Mapped[float]=mapped_column(Float,default=0); gross:Mapped[float]=mapped_column(Float,default=0); ot_hours:Mapped[float]=mapped_column(Float,default=0); ot_rate:Mapped[float]=mapped_column(Float,default=0); ot_amount:Mapped[float]=mapped_column(Float,default=0); absent_days:Mapped[float]=mapped_column(Float,default=0); absent_pct:Mapped[float]=mapped_column(Float,default=0); source_file:Mapped[Optional[str]]=mapped_column(String(255)); __table_args__=(UniqueConstraint('employee_id','salary_month',name='uq_salary_month'),)
class KPITemplate(Base):
    __tablename__='kpi_templates'; id:Mapped[int]=mapped_column(primary_key=True); designation:Mapped[str]=mapped_column(String(160),index=True); period:Mapped[str]=mapped_column(String(20),index=True); title:Mapped[str]=mapped_column(String(255)); description:Mapped[Optional[str]]=mapped_column(Text); weight:Mapped[float]=mapped_column(Float,default=1); max_score:Mapped[float]=mapped_column(Float,default=100); active:Mapped[bool]=mapped_column(Boolean,default=True)
class Evaluation(Base):
    __tablename__='evaluations'; id:Mapped[int]=mapped_column(primary_key=True); employee_id:Mapped[str]=mapped_column(String(64),index=True); kpi_id:Mapped[int]=mapped_column(ForeignKey('kpi_templates.id')); period:Mapped[str]=mapped_column(String(20),index=True); period_key:Mapped[str]=mapped_column(String(30),index=True); answer:Mapped[str]=mapped_column(String(20)); score:Mapped[float]=mapped_column(Float,default=0); note:Mapped[Optional[str]]=mapped_column(Text); evidence_path:Mapped[Optional[str]]=mapped_column(String(500)); evaluated_at:Mapped[datetime]=mapped_column(DateTime,default=datetime.utcnow); evaluator_id:Mapped[str]=mapped_column(String(64))
class Complaint(Base):
    __tablename__='complaints'; id:Mapped[int]=mapped_column(primary_key=True); employee_id:Mapped[str]=mapped_column(String(64),index=True); kind:Mapped[str]=mapped_column(String(30)); subject:Mapped[str]=mapped_column(String(255)); body:Mapped[str]=mapped_column(Text); status:Mapped[str]=mapped_column(String(30),default='OPEN'); attachment:Mapped[Optional[str]]=mapped_column(String(500)); response:Mapped[Optional[str]]=mapped_column(Text); updated_by:Mapped[Optional[str]]=mapped_column(String(64)); created_at:Mapped[datetime]=mapped_column(DateTime,default=datetime.utcnow); updated_at:Mapped[Optional[datetime]]=mapped_column(DateTime)
class Training(Base):
    __tablename__='training_records'; id:Mapped[int]=mapped_column(primary_key=True); employee_id:Mapped[str]=mapped_column(String(64),index=True); training_name:Mapped[str]=mapped_column(String(255)); training_date:Mapped[Optional[date]]=mapped_column(Date); certificate_path:Mapped[Optional[str]]=mapped_column(String(500)); source_file:Mapped[Optional[str]]=mapped_column(String(255))
class Production(Base):
    __tablename__='production_daily'; id:Mapped[int]=mapped_column(primary_key=True); work_date:Mapped[date]=mapped_column(Date,index=True); department:Mapped[str]=mapped_column(String(100)); section:Mapped[Optional[str]]=mapped_column(String(100)); line:Mapped[Optional[str]]=mapped_column(String(50)); style:Mapped[Optional[str]]=mapped_column(String(100)); target:Mapped[float]=mapped_column(Float,default=0); achievement:Mapped[float]=mapped_column(Float,default=0); loss:Mapped[float]=mapped_column(Float,default=0); recovery_plan:Mapped[Optional[str]]=mapped_column(Text); source_file:Mapped[Optional[str]]=mapped_column(String(255))
class Quality(Base):
    __tablename__='quality_records'; id:Mapped[int]=mapped_column(primary_key=True); work_date:Mapped[date]=mapped_column(Date,index=True); section:Mapped[Optional[str]]=mapped_column(String(100)); inspection_qty:Mapped[float]=mapped_column(Float,default=0); defect_qty:Mapped[float]=mapped_column(Float,default=0); reject_qty:Mapped[float]=mapped_column(Float,default=0); dhU:Mapped[float]=mapped_column(Float,default=0); root_cause:Mapped[Optional[str]]=mapped_column(Text); corrective_action:Mapped[Optional[str]]=mapped_column(Text); status:Mapped[str]=mapped_column(String(30),default='OPEN'); source_file:Mapped[Optional[str]]=mapped_column(String(255))
class Location(Base):
    __tablename__='locations'; id:Mapped[int]=mapped_column(primary_key=True); employee_id:Mapped[str]=mapped_column(String(64),index=True); captured_at:Mapped[datetime]=mapped_column(DateTime,default=datetime.utcnow,index=True); latitude:Mapped[float]=mapped_column(Float); longitude:Mapped[float]=mapped_column(Float); accuracy:Mapped[Optional[float]]=mapped_column(Float); consent:Mapped[bool]=mapped_column(Boolean,default=False)
class Document(Base):
    __tablename__='documents'; id:Mapped[int]=mapped_column(primary_key=True); employee_id:Mapped[Optional[str]]=mapped_column(String(64),index=True); category:Mapped[str]=mapped_column(String(80)); name:Mapped[str]=mapped_column(String(255)); path:Mapped[str]=mapped_column(String(500)); uploaded_by:Mapped[str]=mapped_column(String(64)); created_at:Mapped[datetime]=mapped_column(DateTime,default=datetime.utcnow)
class AuditLog(Base):
    __tablename__='audit_logs'; id:Mapped[int]=mapped_column(primary_key=True); actor_id:Mapped[str]=mapped_column(String(64)); action:Mapped[str]=mapped_column(String(120)); entity:Mapped[str]=mapped_column(String(80)); entity_id:Mapped[Optional[str]]=mapped_column(String(80)); created_at:Mapped[datetime]=mapped_column(DateTime,default=datetime.utcnow)
Base.metadata.create_all(engine)
app=FastAPI(title='RIZVI 360° Integrated Workforce Platform',version='2.0.0')
app.add_middleware(CORSMiddleware,allow_origins=['*'],allow_credentials=True,allow_methods=['*'],allow_headers=['*']); app.mount('/uploads',StaticFiles(directory=str(UPLOAD_DIR)),name='uploads')
def db():
    s=SessionLocal()
    try: yield s
    finally: s.close()
def token_for(e): return jwt.encode({'sub':e.employee_id,'role':e.role,'exp':datetime.utcnow()+timedelta(hours=12)},SECRET,algorithm=ALGORITHM)
def current_user(token:str=Depends(oauth2),s:Session=Depends(db)):
    try: p=jwt.decode(token,SECRET,algorithms=[ALGORITHM]); eid=p.get('sub')
    except JWTError: raise HTTPException(401,'Invalid or expired token')
    e=s.query(Employee).filter(Employee.employee_id==eid).first()
    if not e or e.status!='ACTIVE': raise HTTPException(401,'Inactive or missing employee')
    return e
def admin_only(e=Depends(current_user)):
    if e.role not in {'ADMIN','DIRECTOR','GM','HR_ADMIN','DEPT_HEAD'}: raise HTTPException(403,'Admin permission required')
    return e
def save_upload(file, prefix):
    ext=Path(file.filename or '').suffix.lower(); name=f'{prefix}_{secrets.token_hex(8)}{ext}'; path=UPLOAD_DIR/name; return name,path
class KPIIn(BaseModel): designation:str; period:str; title:str; description:Optional[str]=None; weight:float=1; max_score:float=100
class EvalIn(BaseModel): kpi_id:int; period:str; period_key:str; answer:str; score:Optional[float]=None; note:Optional[str]=None
class ComplaintUpdate(BaseModel): status:str; response:Optional[str]=None
class LocationIn(BaseModel): latitude:float; longitude:float; accuracy:Optional[float]=None; consent:bool
class ProductionIn(BaseModel): work_date:date; department:str; section:Optional[str]=None; line:Optional[str]=None; style:Optional[str]=None; target:float=0; achievement:float=0; recovery_plan:Optional[str]=None
class QualityIn(BaseModel): work_date:date; section:Optional[str]=None; inspection_qty:float=0; defect_qty:float=0; reject_qty:float=0; root_cause:Optional[str]=None; corrective_action:Optional[str]=None
@app.get('/api/health')
def health(): return {'ok':True,'service':'RIZVI360','version':'2.0.0'}
@app.post('/api/auth/login')
def login(form:OAuth2PasswordRequestForm=Depends(),s:Session=Depends(db)):
    e=s.query(Employee).filter(or_(Employee.employee_id==form.username,Employee.card_no==form.username)).first()
    if not e or e.status!='ACTIVE' or not pwd.verify(form.password,e.password_hash): raise HTTPException(401,'Invalid ID/card or password')
    return {'access_token':token_for(e),'token_type':'bearer','employee_id':e.employee_id,'role':e.role}
@app.get('/api/me')
def me(e=Depends(current_user)): return {k:v for k,v in e.__dict__.items() if not k.startswith('_') and k!='password_hash'}
def parse_dt(v):
    if not v:return None
    s=str(v).strip();
    for fmt in ('%H:%M:%S','%H:%M'):
        try:return datetime.strptime(s,fmt)
        except: pass
    return None
def daily_hours(p):
    a,b=parse_dt(p.in_time),parse_dt(p.out_time)
    if not a or not b:return 0
    seconds=(b-a).total_seconds()-p.break_minutes*60
    if seconds<0: seconds+=86400
    return max(0,seconds/3600)
def performance_score(s,eid,period=None):
    q=s.query(Evaluation).filter(Evaluation.employee_id==eid); q=q.filter(Evaluation.period==period) if period else q
    vals=[x.score for x in q.all()]; return round(sum(vals)/len(vals),2) if vals else 0
@app.get('/api/dashboard')
def dashboard(e=Depends(current_user),s:Session=Depends(db)):
    punches=s.query(Punch).filter(Punch.employee_id==e.employee_id).all(); hours=sum(daily_hours(p) for p in punches); ot=sum(max(0,daily_hours(p)-(8 if p.punch_date.weekday()!=4 else 0)) for p in punches)
    ev=s.query(Evaluation).filter(Evaluation.employee_id==e.employee_id).all(); avg=sum(x.score for x in ev)/len(ev) if ev else 0; complaints=s.query(Complaint).filter(Complaint.employee_id==e.employee_id).count(); training=s.query(Training).filter(Training.employee_id==e.employee_id).count()
    return {'employee':{'id':e.employee_id,'card_no':e.card_no,'name':e.name,'designation':e.designation,'department':e.department,'section':e.section,'join_date':e.join_date},'attendance':{'punch_days':len(punches),'hours':round(hours,2),'ot_hours':round(ot,2)},'performance':round(avg,2),'complaints':complaints,'trainings':training}
@app.get('/api/dashboard/management')
def management_dashboard(admin=Depends(admin_only),s:Session=Depends(db)):
    active=s.query(Employee).filter(Employee.status=='ACTIVE').count(); total=s.query(Employee).count(); complaints=s.query(Complaint).filter(Complaint.status!='CLOSED').count(); evals=s.query(Evaluation).all(); avg=round(sum(x.score for x in evals)/len(evals),2) if evals else 0; red=sum(1 for x in evals if x.score<60); prod=s.query(Production).all(); target=sum(x.target for x in prod); ach=sum(x.achievement for x in prod); quality=s.query(Quality).all(); insp=sum(x.inspection_qty for x in quality); defects=sum(x.defect_qty for x in quality); return {'employees':{'total':total,'active':active},'performance':{'average':avg,'red_count':red},'complaints':{'open':complaints,'total':s.query(Complaint).count()},'production':{'target':target,'achievement':ach,'achievement_pct':round(ach/target*100,2) if target else 0},'quality':{'inspection':insp,'defects':defects,'dhu':round(defects/insp*100,2) if insp else 0}}
@app.get('/api/kpis')
def kpis(e=Depends(current_user),s:Session=Depends(db),period:str='DAILY'):
    return [{k:v for k,v in x.__dict__.items() if not k.startswith('_')} for x in s.query(KPITemplate).filter(KPITemplate.designation==e.designation,KPITemplate.period==period,KPITemplate.active==True).all()]
@app.post('/api/evaluations')
def evaluate(x:EvalIn,e=Depends(current_user),s:Session=Depends(db)):
    k=s.get(KPITemplate,x.kpi_id)
    if not k or k.designation!=e.designation: raise HTTPException(403,'KPI not assigned to designation')
    ans=x.answer.upper();
    if ans not in {'YES','NO','PARTIAL'}: raise HTTPException(422,'Answer must be YES, NO or PARTIAL')
    if ans=='PARTIAL' and not x.note: raise HTTPException(422,'PARTIAL requires a note')
    score={'YES':k.max_score,'NO':0}.get(ans,max(0,min(k.max_score,float(x.score or 0))))
    ev=Evaluation(employee_id=e.employee_id,kpi_id=k.id,period=x.period,period_key=x.period_key,answer=ans,score=score,note=x.note,evaluator_id=e.employee_id); s.add(ev); s.add(AuditLog(actor_id=e.employee_id,action='CREATE_EVALUATION',entity='evaluation',entity_id='new')); s.commit(); s.refresh(ev); return {'id':ev.id,'score':score,'answer':ans}
@app.get('/api/evaluations/mine')
def my_evaluations(e=Depends(current_user),s:Session=Depends(db)): return [{'id':x.id,'period':x.period,'period_key':x.period_key,'answer':x.answer,'score':x.score,'note':x.note,'evaluated_at':x.evaluated_at} for x in s.query(Evaluation).filter(Evaluation.employee_id==e.employee_id).order_by(Evaluation.evaluated_at.desc()).all()]
@app.post('/api/evaluations/{evaluation_id}/evidence')
async def evidence(evaluation_id:int,file:UploadFile=File(...),e=Depends(current_user),s:Session=Depends(db)):
    ev=s.get(Evaluation,evaluation_id)
    if not ev or ev.employee_id!=e.employee_id: raise HTTPException(404,'Evaluation not found')
    name,path=save_upload(file,f'evidence_{e.employee_id}_{evaluation_id}'); path.write_bytes(await file.read()); ev.evidence_path=f'/uploads/{name}'; s.commit(); return {'path':ev.evidence_path}
@app.post('/api/complaints')
async def complaint(kind:str=Form(...),subject:str=Form(...),body:str=Form(...),file:UploadFile|None=File(None),e=Depends(current_user),s:Session=Depends(db)):
    attachment=None
    if file:
        name,path=save_upload(file,f'case_{e.employee_id}'); path.write_bytes(await file.read()); attachment=f'/uploads/{name}'
    c=Complaint(employee_id=e.employee_id,kind=kind,subject=subject,body=body,attachment=attachment); s.add(c); s.commit(); s.refresh(c); return {'id':c.id,'status':c.status}
@app.get('/api/complaints/mine')
def mine_cases(e=Depends(current_user),s:Session=Depends(db)): return [{'id':x.id,'kind':x.kind,'subject':x.subject,'body':x.body,'status':x.status,'response':x.response,'attachment':x.attachment,'created_at':x.created_at} for x in s.query(Complaint).filter(Complaint.employee_id==e.employee_id).order_by(Complaint.created_at.desc()).all()]
@app.patch('/api/admin/complaints/{case_id}')
def update_case(case_id:int,x:ComplaintUpdate,admin=Depends(admin_only),s:Session=Depends(db)):
    c=s.get(Complaint,case_id)
    if not c: raise HTTPException(404,'Case not found')
    c.status=x.status.upper(); c.response=x.response; c.updated_by=admin.employee_id; c.updated_at=datetime.utcnow(); s.commit(); return {'ok':True}
@app.post('/api/admin/kpis')
def create_kpi(x:KPIIn,admin=Depends(admin_only),s:Session=Depends(db)):
    if x.weight<=0 or x.max_score<=0: raise HTTPException(422,'Weight and max_score must be positive')
    k=KPITemplate(**x.model_dump()); s.add(k); s.commit(); s.refresh(k); return {'id':k.id,'designation':k.designation,'period':k.period,'title':k.title}
@app.post('/api/admin/employees')
def create_employee(data:dict,admin=Depends(admin_only),s:Session=Depends(db)):
    eid=str(data.get('employee_id','')).strip(); card=str(data.get('card_no') or eid); name=str(data.get('name','')).strip()
    if not eid or not name: raise HTTPException(422,'employee_id and name required')
    if s.query(Employee).filter(or_(Employee.employee_id==eid,Employee.card_no==card)).first(): raise HTTPException(409,'Employee exists')
    temp=data.get('password') or secrets.token_urlsafe(8); allowed={c.name for c in Employee.__table__.columns}; obj={k:v for k,v in data.items() if k in allowed and k not in {'id','password_hash'}}; obj.update(employee_id=eid,card_no=card,name=name,password_hash=pwd.hash(temp)); e=Employee(**obj); s.add(e); s.commit(); return {'employee_id':eid,'temporary_password':temp}
def rows_from_file(data,filename):
    if filename.lower().endswith('.csv'): return list(csv.DictReader(io.StringIO(data.decode('utf-8-sig'))))
    from openpyxl import load_workbook
    wb=load_workbook(io.BytesIO(data),read_only=True,data_only=True); ws=wb.active; vals=list(ws.values); headers=[str(x).strip() if x is not None else '' for x in vals[0]]; return [dict(zip(headers,r)) for r in vals[1:]]
def getv(r,*keys):
    for k in keys:
        if k in r and r[k] not in (None,''): return r[k]
    return None
@app.post('/api/admin/import/employees')
async def import_employees(file:UploadFile=File(...),admin=Depends(admin_only),s:Session=Depends(db)):
    rows=rows_from_file(await file.read(),file.filename or 'file.csv'); count=0
    mapping={'designation':('DESIGNATION','Designation'),'department':('DEPARTMENT','Department'),'section':('SECTION','Section'),'floor':('FLOOR','Floor'),'line':('LINE','Line'),'mobile':('MOBILE','Mobile'),'email':('EMAIL','Email'),'grade':('GRADE','Grade'),'category':('CATEGORY','Category'),'manager_id':('MANAGER_ID','Manager ID')}
    for r in rows:
        eid=str(getv(r,'EMPLOYEE_ID','Employee ID','CARD_NO','Card No') or '').strip();
        if not eid: continue
        card=str(getv(r,'CARD_NO','Card No') or eid); e=s.query(Employee).filter(Employee.employee_id==eid).first()
        if not e: e=Employee(employee_id=eid,card_no=card,name=str(getv(r,'NAME','Name') or ''),password_hash=pwd.hash(secrets.token_urlsafe(8))); s.add(e)
        e.card_no=card; e.name=str(getv(r,'NAME','Name') or e.name)
        for a,ks in mapping.items():
            v=getv(r,*ks)
            if v is not None:setattr(e,a,str(v))
        jd=getv(r,'JOIN_DATE','Joining Date','JOINING_DATE');
        if jd:
            try:e.join_date=jd if isinstance(jd,date) else date.fromisoformat(str(jd)[:10])
            except:pass
        count+=1
    s.commit(); return {'imported_or_updated':count}
@app.post('/api/admin/import/punches')
async def import_punches(file:UploadFile=File(...),admin=Depends(admin_only),s:Session=Depends(db)):
    rows=rows_from_file(await file.read(),file.filename or 'file.csv'); count=0
    for r in rows:
        eid=str(getv(r,'EMPLOYEE_ID','Employee ID','CARD_NO','Card No') or '').strip(); ds=getv(r,'DATE','PUNCH_DATE','Punch Date')
        if not eid or not ds: continue
        try:d=ds if isinstance(ds,date) else date.fromisoformat(str(ds)[:10])
        except:continue
        existing=s.query(Punch).filter(Punch.employee_id==eid,Punch.punch_date==d).first(); p=existing or Punch(employee_id=eid,punch_date=d)
        p.in_time=str(getv(r,'IN','IN_TIME','In Time') or '')[:8] or None; p.out_time=str(getv(r,'OUT','OUT_TIME','Out Time') or '')[:8] or None; p.break_minutes=int(float(getv(r,'BREAK_MINUTES','Break Minutes') or 60)); p.source_file=file.filename; s.add(p); count+=1
    s.commit(); return {'imported_or_updated':count}
@app.post('/api/admin/import/salary')
async def import_salary(file:UploadFile=File(...),admin=Depends(admin_only),s:Session=Depends(db)):
    rows=rows_from_file(await file.read(),file.filename or 'file.csv'); count=0
    for r in rows:
        eid=str(getv(r,'EMPLOYEE_ID','Employee ID','CARD_NO','Card No') or '').strip(); month=str(getv(r,'SALARY_MONTH','Month','MONTH') or '')[:7]
        if not eid or not month:continue
        q=s.query(Salary).filter(Salary.employee_id==eid,Salary.salary_month==month).first(); x=q or Salary(employee_id=eid,salary_month=month)
        for a,ks in {'basic':('BASIC','Basic'),'gross':('GROSS','Gross','GROSS_SALARY'),'ot_hours':('OT_HOURS','OT Hour'),'ot_rate':('OT_RATE','OT Rate'),'ot_amount':('OT_AMOUNT','OT Amount'),'absent_days':('ABSENT_DAYS','Absent Days'),'absent_pct':('ABSENT_PCT','Absent %')}.items():
            v=getv(r,*ks)
            if v is not None:
                try:setattr(x,a,float(v))
                except:pass
        x.source_file=file.filename;s.add(x);count+=1
    s.commit();return {'imported_or_updated':count}
@app.post('/api/admin/import/training')
async def import_training(file:UploadFile=File(...),admin=Depends(admin_only),s:Session=Depends(db)):
    rows=rows_from_file(await file.read(),file.filename or 'file.csv'); count=0
    for r in rows:
        eid=str(getv(r,'EMPLOYEE_ID','Employee ID','CARD_NO','Card No') or '').strip(); name=str(getv(r,'TRAINING_NAME','Training','Training Name') or '').strip(); ds=getv(r,'TRAINING_DATE','Date','Training Date')
        if not eid or not name:continue
        d=None
        try:d=ds if isinstance(ds,date) else date.fromisoformat(str(ds)[:10]) if ds else None
        except:pass
        s.add(Training(employee_id=eid,training_name=name,training_date=d,source_file=file.filename));count+=1
    s.commit();return {'imported':count}
@app.post('/api/admin/import/production')
async def import_production(file:UploadFile=File(...),admin=Depends(admin_only),s:Session=Depends(db)):
    rows=rows_from_file(await file.read(),file.filename or 'file.csv');count=0
    for r in rows:
        ds=getv(r,'DATE','WORK_DATE','Date');
        try:d=ds if isinstance(ds,date) else date.fromisoformat(str(ds)[:10])
        except:continue
        target=float(getv(r,'TARGET','Target') or 0);ach=float(getv(r,'ACHIEVEMENT','Achievement','OUTPUT') or 0)
        s.add(Production(work_date=d,department=str(getv(r,'DEPARTMENT','Department') or 'Production'),section=getv(r,'SECTION','Section'),line=getv(r,'LINE','Line'),style=getv(r,'STYLE','Style'),target=target,achievement=ach,loss=max(0,target-ach),recovery_plan=getv(r,'RECOVERY_PLAN','Recovery Plan'),source_file=file.filename));count+=1
    s.commit();return {'imported':count}
@app.post('/api/production')
def add_production(x:ProductionIn,admin=Depends(admin_only),s:Session=Depends(db)):
    p=Production(**x.model_dump(),loss=max(0,x.target-x.achievement));s.add(p);s.commit();return {'id':p.id,'achievement_pct':round(x.achievement/x.target*100,2) if x.target else 0}
@app.post('/api/quality')
def add_quality(x:QualityIn,admin=Depends(admin_only),s:Session=Depends(db)):
    q=Quality(**x.model_dump(),dhU=round(x.defect_qty/x.inspection_qty*100,2) if x.inspection_qty else 0);s.add(q);s.commit();return {'id':q.id,'dhu':q.dhU}
@app.get('/api/admin/training-gap')
def training_gap(admin=Depends(admin_only),s:Session=Depends(db)):
    today=date.today(); result=[]
    for e in s.query(Employee).filter(Employee.status=='ACTIVE').all():
        if not e.join_date:continue
        months=(today.year-e.join_date.year)*12+today.month-e.join_date.month
        if months>=6 and not s.query(Training).filter(Training.employee_id==e.employee_id).first():result.append({'employee_id':e.employee_id,'name':e.name,'join_date':e.join_date,'status':'RED'})
    return result
@app.post('/api/location')
def add_location(x:LocationIn,e=Depends(current_user),s:Session=Depends(db)):
    if not x.consent:raise HTTPException(403,'Location permission/consent is required')
    s.add(Location(employee_id=e.employee_id,**x.model_dump()));s.commit();return {'ok':True}
@app.get('/api/location/mine')
def my_location(e=Depends(current_user),s:Session=Depends(db)):return [{'captured_at':x.captured_at,'latitude':x.latitude,'longitude':x.longitude,'accuracy':x.accuracy} for x in s.query(Location).filter(Location.employee_id==e.employee_id).order_by(Location.captured_at.desc()).limit(100).all()]
@app.post('/api/documents')
async def upload_document(category:str=Form(...),employee_id:Optional[str]=Form(None),file:UploadFile=File(...),admin=Depends(admin_only),s:Session=Depends(db)):
    name,path=save_upload(file,f'doc_{category}');path.write_bytes(await file.read());d=Document(employee_id=employee_id,category=category,name=file.filename or name,path=f'/uploads/{name}',uploaded_by=admin.employee_id);s.add(d);s.commit();return {'id':d.id,'path':d.path}
@app.get('/api/admin/employees')
def employees(admin=Depends(admin_only),s:Session=Depends(db),q:str=''):
    query=s.query(Employee).filter(Employee.status=='ACTIVE');
    if q:query=query.filter(or_(Employee.employee_id.contains(q),Employee.name.contains(q),Employee.department.contains(q),Employee.designation.contains(q)))
    return [{'employee_id':x.employee_id,'card_no':x.card_no,'name':x.name,'designation':x.designation,'department':x.department,'section':x.section,'status':x.status} for x in query.limit(500).all()]

# Real-time, consent-based device location router.
# Imported at the end to avoid disrupting the existing v2 application structure.
from .location_realtime import router as realtime_location_router
app.include_router(realtime_location_router)
